# Prospect asset build: tb.tradez / TB Tradez
# Layout layer follows wiki/playbooks/prospect-asset-doc-standard.md. Fill in, do not redesign.

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

INK    = RGBColor(0x1A, 0x1A, 0x1A)
MUTED  = RGBColor(0x46, 0x4E, 0x5C)
ACCENT = RGBColor(0x4A, 0x53, 0xD9)
RULE   = "D5D5DD"
HEAD   = "F2F3F7"
BODY, MONO = "Georgia", "Consolas"

SHOTS = r"C:\Users\Lenovoi\Desktop\Prospect Assets\tb-tradez"
REFS  = r"C:\Users\Lenovoi\Downloads\Pulsify-Brain\raw-sources\inspiration\thumbnail-retouch-references"
OUT   = os.path.join(SHOTS, "tbtradez-four-videos-one-picture.docx")


def set_font(run, name=BODY, size=10.5, color=INK, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold, run.italic = bold, italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)


def spacing(p, before=0, after=8, line=1.25):
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line


def para(doc, text, size=10.5, color=INK, after=8, before=0, bold=False,
         italic=False, name=BODY, indent=0, line=1.25):
    p = doc.add_paragraph()
    spacing(p, before, after, line)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    set_font(p.add_run(text), name, size, color, bold, italic)
    return p


def bottom_rule(p, color=RULE, sz=6):
    ppr = p._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "4"); b.set(qn("w:color"), color)
    bdr.append(b); ppr.append(bdr)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def cell_border(cell, color=RULE, sz=4):
    tcpr = cell._tc.get_or_add_tcPr()
    bs = OxmlElement("w:tcBorders"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "0"); b.set(qn("w:color"), color)
    bs.append(b); tcpr.append(bs)


def no_borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    bs = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "none")
        bs.append(e)
    pr.append(bs)


def heading(doc, text):
    return para(doc, text, size=12.5, bold=True, before=15, after=3)


def rebuild(doc, num, title_text, direction_text):
    p = doc.add_paragraph()
    spacing(p, 9, 1, 1.2)
    set_font(p.add_run(f"{num}.  "), BODY, 10.5, ACCENT, bold=True)
    set_font(p.add_run(title_text), MONO, 9.5, INK)
    para(doc, direction_text, size=9, color=MUTED, after=8, indent=0.22, line=1.2)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(1.2)
sec.top_margin = sec.bottom_margin = Inches(0.75)

normal = doc.styles["Normal"]
normal.font.name, normal.font.size, normal.font.color.rgb = BODY, Pt(10.5), INK
rpr = normal.element.get_or_add_rPr(); rf = OxmlElement("w:rFonts")
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), BODY)
rpr.append(rf)

# --- title ----------------------------------------------------------------
para(doc, "Four videos, one picture", size=16, bold=True, after=2, line=1.1)
sub = para(doc, "Built for TB from his own uploads, his reels, his transcripts and tbtradez.com. "
                "August 2026.", size=9, color=MUTED, after=13, line=1.2)
bottom_rule(sub)

# --- the finding ----------------------------------------------------------
heading(doc, "The finding")
para(doc, "One of your reels did 101,000 views. You have 2,243 followers. Reach is not your "
          "problem and it has not been for a while. Nothing sticking to it is.")
para(doc, "The reel that did it was the beginner roadmap: pick a strategy, backtest it, journal, "
          "stay disciplined. A hundred thousand people watched that and almost none of them "
          "followed, because everything of value in it was delivered in thirty seconds and none "
          "of it was yours. The video that would have made them follow, the 140 point NQ read, "
          "did 1,141.", after=14)

# --- evidence table -------------------------------------------------------
heading(doc, "Your last four uploads")

CONTENT = Inches(6.1)
LEFT_W, RIGHT_W = Inches(3.85), Inches(2.25)

rows = [
    ("Liquidity Explained in Under 9 Mins",
     "0 views. The headline runs behind your hood, so the one word carrying it gets cut.",
     "yt-1FlbO1Xm4xE.jpg"),
    ("Market Structure",
     "69 views. Same photograph as the one above.",
     "yt-xyJ2LMZTldc.jpg"),
    ("How to Read Candlesticks",
     "220 views. Same photograph again.",
     "yt-mVm6Dhz21hg.jpg"),
    ("Why 90% Fail",
     "133 views. Same photograph again.",
     "yt-ouT5JtOJHQs.jpg"),
]

t = doc.add_table(rows=len(rows) + 1, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.LEFT
t.autofit = False

hdr = t.rows[0]
for c, label in zip(hdr.cells, ("Your title", "What every one of them looks like")):
    c.text = ""
    p = c.paragraphs[0]; spacing(p, 3, 3, 1.15)
    set_font(p.add_run(label), BODY, 9, MUTED, bold=True)
    shade(c, HEAD); cell_border(c)
hdr.cells[0].width, hdr.cells[1].width = LEFT_W, RIGHT_W

for i, (title_text, note, img) in enumerate(rows, start=1):
    r = t.rows[i]
    left, right = r.cells[0], r.cells[1]
    left.text = ""
    p1 = left.paragraphs[0]; spacing(p1, 5, 2, 1.2)
    set_font(p1.add_run(title_text), BODY, 9.5, INK, bold=True)
    p2 = left.add_paragraph(); spacing(p2, 0, 5, 1.2)
    set_font(p2.add_run(note), BODY, 8.5, MUTED, italic=True)
    right.text = ""
    p3 = right.paragraphs[0]; spacing(p3, 4, 4, 1.0)
    p3.add_run().add_picture(os.path.join(SHOTS, img), width=Inches(1.95))
    left.width, right.width = LEFT_W, RIGHT_W
    cell_border(left); cell_border(right)

para(doc, "Nobody who clicked the first one has a reason to click the second. On a channel with "
          "seven videos, that is most of the catalogue looking like the same upload posted four "
          "times.", size=9, color=MUTED, before=7, after=4)

# --- the clock ------------------------------------------------------------
heading(doc, "The clock says 1:89")
para(doc, "Look at the phone you are holding. The time in the status bar reads 1:89, on all four "
          "of them. There is no such time. The screen in your hand is made up, and it is the "
          "same made-up screen every time.")
para(doc, "It gets worse when the numbers are put side by side. Those four say TRADING "
          "$1,600,000. The Bali video shows a portfolio at $27,540.82, up 8.47% on the day. Same "
          "channel, same person, two figures that do not agree with each other.")
para(doc, "This matters more in your niche than almost any other, because a payout number on a "
          "thumbnail is the one piece of proof everybody knows can be faked. Your audience has "
          "scrolled past thousands of them. A viewer who spots the clock does not think you made "
          "a design mistake, and once that thought lands it is very hard to unland.")
para(doc, "The honest part of those images is your face. No retouching, real skin, you look "
          "nineteen. That is the asset. Everything you are holding is working against it, and "
          "your own numbers agree: the reel that reached a hundred thousand people has no money "
          "in the frame at all. Just you, a hoodie, talking.")

# --- what is working ------------------------------------------------------
heading(doc, "What is already working")
para(doc, "You are shipping long form consistently, which is the hard part. Seventeen minutes on "
          "market structure, thirteen on candlesticks, in a two week window. Most people talking "
          "about starting a channel have not uploaded twice.")
para(doc, "And the 140 point read is genuinely good. You did not just call the trade, you "
          "explained why you took a setup other people were fading: there was a 30 minute "
          "unmitigated gap sitting below your 15 minute one, and your read was that this market "
          "has not manipulated into a single 30 minute gap all cycle. That is a real opinion "
          "backed by something you have actually watched. It is also the only thing in your "
          "whole library that another trader could not have said.")

# --- whose method -------------------------------------------------------
heading(doc, "Whose method is it")
para(doc, "Liquidity. Market structure. Candlesticks. Those are ICT concepts, and hundreds of "
          "channels already own them. Your site says ICT strategy. Your reel closes with \"follow "
          "to learn everything I know about ICT.\" So the channel is currently building an "
          "audience for somebody else's method, and every one of those videos is competing "
          "against people who have taught the same concept for years.")
para(doc, "Meanwhile your own idea is sitting at the end of a reel with nothing attached to it. "
          "Think in objectives. Do not ask where your stop is, ask what objectives exist at your "
          "stop against what exists at your target, and wait for the higher timeframe one to get "
          "taken first. That is teachable, it is nameable, and as far as your channel is "
          "concerned it does not exist yet.")
para(doc, "It is also the difference between a channel people watch and a thing people pay to be "
          "taught. Nobody pays for liquidity explained. They pay for the specific way you see it "
          "that they cannot get anywhere else.")

# --- the rebuild ----------------------------------------------------------
heading(doc, "Six you could make")
para(doc, "Every one of these comes from something you have already said on camera or already "
          "have on your account. Nothing needs a new idea.",
     size=9, color=MUTED, after=6)

rebuild(doc, 1, "think in objectives (the fix for getting stop hunted)",
        "Thumbnail: objective at your stop. Your chart with the higher timeframe gap marked "
        "below the stop. Your idea, named, first. Everything else on the channel should point "
        "back at this one.")
rebuild(doc, 2, "how i caught 140 points on nq before the open (step-by-step)",
        "Thumbnail: 9:31, one entry. The actual chart with your entry marked, face small in the "
        "corner. No phone, no number in your hand.")
rebuild(doc, 3, "the 30 minute gap nobody has traded all cycle (full breakdown)",
        "Thumbnail: nobody traded this. The gap highlighted on your own chart and nothing else "
        "in frame. This is the strongest opinion you have voiced and it is currently thirty "
        "seconds long.")
rebuild(doc, 4, "$70 away from passing an eval, then this (full breakdown)",
        "Thumbnail: $70 short. The real eval screen, your own. A number you can actually stand "
        "behind beats one you cannot, every time.")
rebuild(doc, 5, "why 90% fail, and it isn't discipline (full breakdown)",
        "Thumbnail: not discipline. You at the desk, hood down, nothing in your hands. Same "
        "topic you already published, with a turn on it and a picture that is not the other four.")
rebuild(doc, 6, "the 5 rules i'd give myself at 17 (step-by-step)",
        "Thumbnail: before my first eval. Your reach audience already told you they want this, "
        "and being nineteen is the part of it nobody else can copy.")

para(doc, "The four already published do not need deleting. They need four different pictures "
          "and titles without \"TB Trading Concept\" on the end, which is eating the characters "
          "that show up in search and adding nothing.", size=9, color=MUTED, before=8, after=4)

# --- references -----------------------------------------------------------
heading(doc, "What this looks like elsewhere")
para(doc, "Three from other channels, untouched. The same two moves in all of them: the proof is "
          "the picture, and the text stays out of the way.",
     size=9, color=MUTED, after=8)

ref_rows = [
    ("XQB7y00msmc-HD.jpg",
     "Screen recording. The figures are the biggest thing in the frame and the face sits in the "
     "corner. No headline, no arrows, nothing added."),
    ("3QDKzi4nFl8-HD.jpg",
     "Real desk, real window, two real people mid-work. One short line across the top and a "
     "label on each of them. Nothing staged, no lifestyle in shot."),
    ("KnbqLjGzD3k-HD.jpg",
     "Same two traders, and no numbers anywhere. The question the video answers is just written "
     "on the picture. Quiet works in this niche."),
]

REF_IMG_W, REF_TXT_W = Inches(2.7), Inches(3.4)
rt = doc.add_table(rows=len(ref_rows), cols=2)
rt.alignment = WD_TABLE_ALIGNMENT.LEFT
rt.autofit = False
no_borders(rt)
for i, (f, cap) in enumerate(ref_rows):
    r = rt.rows[i]
    lc, rc = r.cells[0], r.cells[1]
    lc.text = ""
    p1 = lc.paragraphs[0]; spacing(p1, 5, 5, 1.0)
    p1.add_run().add_picture(os.path.join(REFS, f), width=Inches(2.55))
    rc.text = ""
    p2 = rc.paragraphs[0]; spacing(p2, 5, 5, 1.2)
    set_font(p2.add_run(cap), BODY, 8.5, MUTED)
    lc.width, rc.width = REF_IMG_W, REF_TXT_W

# --- the site -------------------------------------------------------------
heading(doc, "tbtradez.com")
para(doc, "Everything on that page is a transaction. A discount code, a giveaway, a free room. "
          "The money mechanism underneath it is prop firm affiliate commission, which pays you "
          "once when somebody signs up whether or not they ever get anywhere. That is real "
          "revenue and it is worth keeping, but it is rented, it caps out, and it means your "
          "incentive stops at the signup rather than at whether the person succeeds.")
para(doc, "Nothing on the page carries your actual thinking. Not the objectives idea, not the 30 "
          "minute gap read, none of it. Somebody arrives from a reel, gets offered a code and a "
          "Discord invite, and leaves knowing nothing about how you see the market.")
para(doc, "Four of the six prop firm cards say coming soon with no date, which reads as "
          "unfinished rather than upcoming. The 31 Discord screenshots are rendered twice in the "
          "markup, so the page repeats itself. The X icon has no link behind it. And the first "
          "question in your FAQ is whether this is legitimate, which is worth sitting with: if "
          "the page has to answer that, something upstream raised it.")

# --- what it looks like built --------------------------------------------
heading(doc, "What it looks like once it is joined up")
para(doc, "Right now you have two accounts sharing one handle. The broad stuff reaches a hundred "
          "thousand people and converts nobody. The deep stuff would convert and reaches "
          "eleven hundred. Nothing runs between them.")
para(doc, "Joined up, the beginner reel is the front door rather than the whole house. Somebody "
          "watches it, finds one line in it that only you would say, follows that to a video "
          "where you spend fifteen minutes proving it, and arrives in the Discord already sold "
          "on how you think. Same content you are making now, pointed at each other.")
para(doc, "What that changes is who turns up. Today anybody who messages you starts from zero "
          "and you end up proving yourself one conversation at a time. When somebody has watched "
          "you work through three trades first, the proving is done before they type, and the "
          "conversation is about whether they are a fit rather than whether you are real. Those "
          "are the only people worth spending an hour on.")

# --- horizon note ---------------------------------------------------------
heading(doc, "One last thing")
para(doc, "You reached 101,000 people on 2,243 followers. Almost nobody gets that kind of reach "
          "at your size. Whatever is wrong here, it is not the follower count, and chasing that "
          "number is the wrong thing to spend the next six months on.")
para(doc, "What is above is foundation work, and foundations are slow to show in the numbers. "
          "That is normal rather than a sign it is not working. Most people quit somewhere "
          "around week three, which is usually right before it starts to compound.", after=0)

doc.save(OUT)
print("saved:", OUT)
print("bytes:", os.path.getsize(OUT))
