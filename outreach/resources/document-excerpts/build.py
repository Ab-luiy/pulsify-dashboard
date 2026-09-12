from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

ROOT = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.72)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(20,20,20)
normal.paragraph_format.line_spacing = 1.14
normal.paragraph_format.space_after = Pt(9)
for name, size in [('Title',23), ('Heading 1',17), ('Heading 2',12)]:
    s=doc.styles[name]
    s.font.name='Arial'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    s.font.bold = name != 'Title'
    s.paragraph_format.space_before=Pt(11)
    s.paragraph_format.space_after=Pt(8)
    s.paragraph_format.keep_with_next=True
doc.core_properties.author='Pulsify'
doc.core_properties.title='Content profile and VSL working notes'

def p(text='', bold=False, highlight=False, italic=False):
    para=doc.add_paragraph()
    r=para.add_run(text); r.bold=bold; r.italic=italic
    if highlight: r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    return para
def h(text,level=2): doc.add_heading(text,level)
def quote(text):
    para=p(text)
    para.paragraph_format.left_indent=Inches(.18)
    return para
def page(title,intro):
    if len(doc.paragraphs): doc.add_page_break()
    doc.add_paragraph(title,'Title')
    p(intro)

page('Content planning notes',
     'When planning the next recording, start with the questions and decisions already coming up in your work. A useful reply can become the outline.')
h('When you know your subject but have no video idea',1)
p('One real question can give you several video angles.',bold=True,highlight=True)
h('Start with a question you have answered recently')
p('Check your comments, DMs or notes from a call. Pick something your audience actually asks, then find one example you can explain or show.')
h('Turn that question into three directions')
quote('The answer: how you approach the problem, step by step.\nThe mistake: what people overlook when trying to solve it.\nThe decision: how you choose between two possible approaches.')
h('For example')
p('If the question is "how do you choose what to test?", you could make:')
quote('1. What I check before testing something new\n2. The mistake that makes a test hard to learn from\n3. Why I chose this idea over the other one')
p('The example could be a trade setup, a product or a video angle. Use something you actually worked through so you can show your reasoning.')
h('Before you record')
p('Choose one direction. Write the opening in your own words, pick the example you will show, and decide what the viewer should understand by the end.')
p('Packaging check: the title sets the expectation; the opening and example need to deliver it. Your thumbnail should add information rather than repeat the title.',italic=True)

page('Instagram profile notes',
     'The bio, pinned posts and highlights work together. Each has a role in helping a new visitor understand your subject, your credibility and the next step.')
h('Give someone new a place to start',1)
p('The first highlight is the introduction to the rest of the profile.',bold=True,highlight=True)
h('Build the start here sequence')
p('Use this ten-slide starting structure. The label can match your voice; the job of the sequence stays the same.')
quote('Slides 1-3: your story and the problem you came to understand.\nSlides 4-6: student proof with the starting point and outcome.\nSlides 7-8: how your approach works.\nSlides 9-10: the next step and where to find it.')
h('Give the other highlights a clear job')
p('Keep proof, your approach and the application path easy to find. A proof highlight needs context around the result. A highlight about your approach needs to explain how you work.')
h('Make the bio agree with that journey')
p('A new visitor should be able to tell what they will get from your content, why they should trust you, and what to do next. Use proof they can trace back to its source.')
p('Keep a route into your YouTube content for people who want to understand more, and a clear DM route for people ready to talk.',highlight=True)
p('If the bio asks for a keyword, the reply needs to deliver what the bio offered. Check that the words, link and actual destination agree.')
h('Review before changing the labels')
p('Check what is inside each highlight against its job. Changing the cover or name alone does not build the missing introduction, explanation or proof.')

page('VSL working notes',
     'The page and video need to continue the promise that brought someone there. Review the message before choosing a new layout or adding more sections.')
h('Make the first minute useful on its own',1)
p('Use the first 60-90 seconds as a short version of the full VSL.',bold=True,highlight=True)
p('Give the viewer enough context to understand who it is for, what you help with, why you can help, how it works and what happens next. Expand those points through the rest of the video.')
h('Check the order of the full explanation')
quote('1. Hook: what you can help them accomplish.\n2. Why you: relevant experience and believable proof.\n3. Buying motives: the reasons people want the solution.\n4. Offer: what they receive and how it works.\n5. Objections: the concerns that need answering.\n6. Qualification and CTA: who it fits and the next step.')
h('Check the handoff from content to page')
p('Read the message that brings someone in, then the headline, then the video opening. They should describe the same problem and promise. Any change in meaning needs fixing before the design pass.')
h('Use the questions people actually ask')
p('Pull the explanation and objection sections from real conversations. If several people need the same clarification on a call, check whether the page or video answers it clearly.')
h('Find the point to investigate')
p('Separate page visits, video starts, application starts, qualified bookings and calls attended. Compare the stages to locate the drop-off before deciding which part to change.')
p('This is a structure check. Performance claims need the actual funnel data.',italic=True)

doc.save(ROOT/'pulsify-working-notes.docx')
print(ROOT/'pulsify-working-notes.docx')
