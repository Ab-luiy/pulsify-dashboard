"""Pre-send checks on a prospect asset docx.

    python tools/check_prospect_doc.py "C:\\...\\creator-doc.docx"

Checks what can be checked without a renderer: voice violations, unfilled
placeholders, image count, word count against the length rule, and whether the
document is actually a valid docx.

This does NOT replace looking at the rendered pages. Nothing ships unlooked at.
"""

import os
import re
import sys
import zipfile

BANNED = {
    "em dash": "\u2014",
    "en dash": "\u2013",
    "curly apostrophe": "\u2019",
    "curly quote open": "\u201c",
    "curly quote close": "\u201d",
    "arrow glyph": "\u2192",
    "ellipsis glyph": "\u2026",
}

TELLS = [
    (r"\bit'?s not\b[^.]{0,60}\bit'?s\b", "negative parallelism"),
    (r"\bhere'?s the (thing|kicker)\b", "false suspense"),
    (r"\bserves as\b|\brepresents a\b|\bstands as\b", "copula dodge"),
    (r"\bit'?s worth noting\b|\bimportantly\b|\bnotably\b", "filler"),
    (r"\blet'?s (unpack|break this down|dive in)\b", "false suspense"),
    (r"\bin conclusion\b|\bin today'?s\b", "AI boilerplate"),
]

WORD_CEILING = 1500  # two pages of prose plus image tables


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    path = sys.argv[1]
    if not os.path.exists(path):
        print(f"not found: {path}")
        return 1

    try:
        import docx
    except ImportError:
        print("pip install python-docx")
        return 1

    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    text = "\n".join(parts)
    words = len(text.split())

    with zipfile.ZipFile(path) as z:
        images = [n for n in z.namelist() if n.startswith("word/media/")]

    fails = []

    print(f"file    {os.path.basename(path)}  ({os.path.getsize(path):,} bytes)")
    print(f"words   {words}")
    print(f"tables  {len(d.tables)}   images {len(images)}")
    print()

    if words > WORD_CEILING:
        fails.append(f"{words} words, over the {WORD_CEILING} ceiling. Cut a section, "
                     f"do not trim sentences.")

    print("voice")
    for label, ch in BANNED.items():
        n = text.count(ch)
        print(f"  {label:<20} {n}")
        if n:
            fails.append(f"{n} {label}")

    print("\ntells")
    for pattern, label in TELLS:
        hits = re.findall(pattern, text, re.I)
        if hits:
            print(f"  {label:<20} {len(hits)}  {hits[:2]}")
            fails.append(f"{label} x{len(hits)}")
    if not any(re.findall(p, text, re.I) for p, _ in TELLS):
        print("  clean")

    print("\nplaceholders")
    stray = re.findall(r"TODO|TBC|XXX|lorem|PLACEHOLDER", text, re.I)
    slots = re.findall(r"\[[^\]]+\]", text)
    print(f"  stray      {len(stray)}")
    print(f"  fill slots {slots if slots else 'none'}")
    if stray:
        fails.append(f"{len(stray)} stray placeholders")

    print()
    if fails:
        print("FAILED")
        for f in fails:
            print(f"  - {f}")
    else:
        print("PASSED the automated checks.")
    print("\nStill required: render to PDF, rasterise the pages, and look at every one.")
    return 1 if fails else 0


if __name__ == "__main__":
    raise SystemExit(main())
